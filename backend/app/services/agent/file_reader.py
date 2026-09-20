from __future__ import annotations

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 50000

_TEXT_MIME_TYPES = {
    "text/plain", "text/csv", "text/markdown", "text/html",
    "application/json", "application/xml", "text/xml",
}

_PDF_MIME_TYPE = "application/pdf"

_DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def extract_file_text(media_type: str | None, content: bytes) -> dict:
    if len(content) > MAX_FILE_BYTES:
        return {
            "text": f"(file too large: {len(content)} bytes, max {MAX_FILE_BYTES})",
            "truncated": False,
        }

    text = None
    truncated = False

    if media_type in _TEXT_MIME_TYPES:
        text = _extract_text_encoded(content)
    elif media_type == _PDF_MIME_TYPE:
        text = _extract_pdf(content)
    elif media_type == _DOCX_MIME_TYPE:
        text = _extract_docx(content)
    elif media_type == _XLSX_MIME_TYPE:
        text = _extract_xlsx(content)
    elif media_type == _PPTX_MIME_TYPE:
        text = _extract_pptx(content)

    if text is None:
        return {
            "text": f"(binary file, type: {media_type or 'unknown'}, size: {len(content)} bytes, no text extracted)",
            "truncated": False,
            "error": True,
        }

    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        truncated = True

    return {"text": text, "truncated": truncated}


def _extract_text_encoded(content: bytes) -> str | None:
    for encoding in ("utf-8", "latin-1", "gbk"):
        try:
            decoded = content.decode(encoding)
            if "\x00" in decoded:
                continue
            return decoded
        except (UnicodeDecodeError, LookupError):
            continue
    return None


def _extract_pdf(content: bytes) -> str | None:
    try:
        import io
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
        return "\n".join(parts) if parts else None
    except Exception:
        return None


def _extract_docx(content: bytes) -> str | None:
    try:
        import io
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                if row_text.strip():
                    parts.append(row_text)
        return "\n".join(parts) if parts else None
    except Exception:
        return None


def _extract_xlsx(content: bytes) -> str | None:
    try:
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"--- Sheet: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                row_values = [str(cell) if cell is not None else "" for cell in row]
                if any(v.strip() for v in row_values):
                    parts.append(" | ".join(row_values))
        wb.close()
        return "\n".join(parts) if parts else None
    except Exception:
        return None


def _extract_pptx(content: bytes) -> str | None:
    try:
        import io
        from pptx import Presentation
        prs = Presentation(io.BytesIO(content))
        parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            parts.append(f"--- Slide {slide_num} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            parts.append(text)
        return "\n".join(parts) if parts else None
    except Exception:
        return None
